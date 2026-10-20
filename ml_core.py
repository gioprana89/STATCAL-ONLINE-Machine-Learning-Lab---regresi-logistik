"""Logika regresi logistik dan laporan Excel untuk STATCAL ML Lab.

Kode ini terpisah dari antarmuka Streamlit supaya mudah diuji.
Seluruh angka evaluasi dihitung pada DATA TESTING (bukan training).
"""
from __future__ import annotations

from io import BytesIO
import numpy as np
import pandas as pd
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score, confusion_matrix, f1_score, precision_score,
    recall_score, roc_auc_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler

ID_COLUMNS = {"ID_Mahasiswa"}
DEFAULT_TARGET = "Kelulusan_MK"
GUIDE = {
    "ID_Mahasiswa": "Pengenal baris mahasiswa; tidak digunakan untuk pelatihan.",
    "Jam_Belajar_Mingguan": "Lama belajar per minggu.",
    "Kehadiran_Persen": "Persentase kehadiran dalam pembelajaran.",
    "Nilai_Tugas_Awal": "Nilai tugas yang diketahui sebelum akhir semester.",
    "Nilai_Kuis_Awal": "Nilai kuis yang diketahui sebelum akhir semester.",
    "Aktivitas_Belajar": "Indeks aktivitas belajar.",
    "Akses_LMS_Mingguan": "Frekuensi akses LMS per minggu.",
    "Tugas_Tepat_Waktu_Persen": "Persentase tugas diserahkan tepat waktu.",
    "Metode_Belajar": "Metode belajar: Mandiri, Kelompok, atau Campuran.",
    "Kelulusan_MK": "Status akhir mata kuliah; variabel target (Y).",
}


def build_pipeline(numeric_cols, categorical_cols, C=1.0, balanced=False):
    blocks = []
    if numeric_cols:
        blocks.append(("numerik", Pipeline([
            ("isi_kosong", SimpleImputer(strategy="median")),
            ("standar", StandardScaler()),
        ]), numeric_cols))
    if categorical_cols:
        blocks.append(("kategori", Pipeline([
            ("isi_kosong", SimpleImputer(strategy="most_frequent")),
            ("one_hot", OneHotEncoder(handle_unknown="ignore", drop="first", sparse_output=False)),
        ]), categorical_cols))
    preprocess = ColumnTransformer(blocks, remainder="drop")
    return Pipeline([
        ("persiapan", preprocess),
        ("regresi_logistik", LogisticRegression(
            C=float(C), solver="lbfgs", max_iter=2000,
            class_weight="balanced" if balanced else None,
        )),
    ])


def run_logistic_analysis(
    raw_df: pd.DataFrame, target: str, positive: str,
    features: list[str], train_percent: int = 80,
    threshold: float = 0.5, C: float = 1.0,
    balanced: bool = False, random_seed: int = 42,
) -> dict:
    if raw_df.empty or not features:
        raise ValueError("Dataset harus berisi data, dan sedikitnya satu variabel X dipilih.")
    if target not in raw_df.columns or any(c not in raw_df.columns for c in features):
        raise ValueError("Variabel X atau Y tidak ditemukan di dataset.")
    if target in features or any(c in ID_COLUMNS for c in features):
        raise ValueError("Variabel target dan ID tidak boleh digunakan sebagai prediktor.")
    if not 0 < train_percent < 100 or not 0 < threshold < 1:
        raise ValueError("Persentase training dan threshold tidak valid.")
    df = raw_df.loc[raw_df[target].notna()].copy().reset_index(drop=True)
    classes = sorted(df[target].astype(str).unique().tolist())
    if len(classes) != 2 or positive not in classes:
        raise ValueError("Y harus memiliki tepat dua kategori dan kelas positif harus valid.")
    negative = next(c for c in classes if c != positive)
    X = df[features].copy()
    for col in features:
        if not pd.api.types.is_numeric_dtype(X[col]):
            X[col] = X[col].astype("string").replace(r"^\s*$", pd.NA, regex=True).astype(object)
    if any(X[col].isna().all() for col in features):
        raise ValueError("Ada fitur X yang seluruh nilainya kosong. Hapus fitur tersebut.")
    numeric_cols = [c for c in features if pd.api.types.is_numeric_dtype(X[c])]
    categorical_cols = [c for c in features if c not in numeric_cols]
    y = (df[target].astype(str) == positive).astype(int)
    if min(y.value_counts()) < 2:
        raise ValueError("Masing-masing kelas Y harus memiliki sedikitnya dua observasi.")
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=(100 - train_percent) / 100,
        random_state=int(random_seed), stratify=y,
    )
    model = build_pipeline(numeric_cols, categorical_cols, C, balanced)
    model.fit(X_train, y_train)
    scores = model.predict_proba(X_test)[:, 1]
    predicted = (scores >= threshold).astype(int)
    cm = confusion_matrix(y_test, predicted, labels=[0, 1])
    tn, fp, fn, tp = map(int, cm.ravel())
    recall = float(recall_score(y_test, predicted, zero_division=0))
    metrics = {
        "Accuracy": float(accuracy_score(y_test, predicted)),
        "Precision": float(precision_score(y_test, predicted, zero_division=0)),
        "Recall": recall,
        "F1-score": float(f1_score(y_test, predicted, zero_division=0)),
        "ROC-AUC": float(roc_auc_score(y_test, scores)) if y_test.nunique() == 2 else None,
        "TN": tn, "FP": fp, "FN": fn, "TP": tp,
        "Jumlah Aktual Positif": tp + fn,
        "Jumlah Prediksi Positif": tp + fp,
        "N Training": len(X_train), "N Testing": len(X_test),
    }
    base_class = int(y_train.mean() >= 0.5)
    metrics["Accuracy Baseline Mayoritas"] = float(accuracy_score(
        y_test, np.full(len(y_test), base_class)))
    category = np.select(
        [(y_test.to_numpy()==0)&(predicted==0),
         (y_test.to_numpy()==0)&(predicted==1),
         (y_test.to_numpy()==1)&(predicted==0),
         (y_test.to_numpy()==1)&(predicted==1)],
        ["TN", "FP", "FN", "TP"], default="-",
    )
    full = df.copy()
    full.insert(0, "Nomor_Baris_Data", np.arange(1, len(full) + 1))
    full["Pembagian_Data"] = ""
    full.loc[X_train.index, "Pembagian_Data"] = "Training"
    full.loc[X_test.index, "Pembagian_Data"] = "Testing"
    training = full.loc[X_train.index].sort_values("Nomor_Baris_Data").reset_index(drop=True)
    testing = full.loc[X_test.index].copy()
    testing["Kelas_Aktual"] = np.where(y_test.to_numpy() == 1, positive, negative)
    testing["Kode_Aktual_0_1"] = y_test.to_numpy().astype(int)
    testing["Probabilitas_Kelas_Positif"] = scores
    testing["Probabilitas_Kelas_Negatif"] = 1.0 - scores
    testing["Kelas_Prediksi"] = np.where(predicted == 1, positive, negative)
    testing["Kode_Prediksi_0_1"] = predicted
    testing["Prediksi_Benar"] = np.where(y_test.to_numpy() == predicted, "Ya", "Tidak")
    testing["Kategori_Confusion_Matrix"] = category
    testing = testing.sort_values("Nomor_Baris_Data").reset_index(drop=True)
    cols = (["Nomor_Baris_Data"] + [c for c in ["ID_Mahasiswa"] if c in df.columns] +
            ["Kelas_Aktual", "Kelas_Prediksi", "Probabilitas_Kelas_Positif",
             "Probabilitas_Kelas_Negatif", "Prediksi_Benar", "Kategori_Confusion_Matrix"])
    prediction = testing[cols].copy()
    return {
        "df": df, "full": full, "training": training, "testing": testing,
        "prediction": prediction, "model": model,
        "X_train": X_train, "X_test": X_test,
        "y_train": y_train, "y_test": y_test,
        "scores": scores, "predicted": predicted,
        "cm": cm, "metrics": metrics, "negative": negative, "positive": positive,
        "numeric": numeric_cols, "categorical": categorical_cols,
        "config": {
            "Target Y": target, "Kelas positif (kode 1)": positive,
            "Kelas negatif (kode 0)": negative,
            "Prediktor X": ", ".join(features),
            "Training (%)": train_percent, "Testing (%)": 100-train_percent,
            "Threshold": threshold, "Regularisasi C": C,
            "Bobot kelas seimbang": "Ya" if balanced else "Tidak",
            "Random seed": random_seed, "Jumlah observasi": len(df),
        },
    }


def predict_student_profile(result: dict, entry: dict) -> dict:
    features = result["X_train"].columns.tolist()
    profile_df = pd.DataFrame([entry], columns=features)
    probability = float(result["model"].predict_proba(profile_df)[0, 1])
    threshold = result["config"]["Threshold"]
    return {
        "input": dict(entry), "probability": probability,
        "prediction": result["positive"] if probability >= threshold else result["negative"],
        "threshold": threshold,
    }


def create_excel_report(result: dict, profile: dict | None = None) -> bytes:
    """Workbook hasil aplikasi; semua nilai dibuat dari hasil run yang sama."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    m, cfg = result["metrics"], result["config"]
    pos, neg = result["positive"], result["negative"]
    questions = [
        (1, f"Berapa mahasiswa yang diprediksi {pos} dengan benar?",
         "True Positive (TP)", m["TP"],
         f"{m['TP']} mahasiswa pada data testing sebenarnya {pos} dan diprediksi {pos}."),
        (2, f"Berapa mahasiswa yang sebenarnya {pos} tetapi keliru diprediksi {neg}?",
         "False Negative (FN)", m["FN"],
         f"{m['FN']} mahasiswa pada data testing sebenarnya {pos} tetapi diprediksi {neg}."),
        (3, f"Seberapa baik model mengenali mahasiswa berisiko {pos}?",
         "Recall", m["Recall"],
         f"{m['TP']} dari {m['TP']+m['FN']} mahasiswa yang sebenarnya {pos} berhasil dikenali (Recall {m['Recall']:.1%})."),
        (4, f"Berapa probabilitas seorang mahasiswa diprediksi {pos}?",
         "Probabilitas profil", profile["probability"] if profile else None,
         (f"Profil simulasi menghasilkan probabilitas {pos} sebesar {profile['probability']:.1%}; kelas prediksi: {profile['prediction']}."
          if profile else "Belum tersedia: jalankan simulasi profil pada Tab 4 sebelum mengunduh ulang Excel.")),
    ]
    questions_df = pd.DataFrame(questions, columns=[
        "No", "Pertanyaan", "Indikator", "Hasil_Numerik", "Interpretasi",
    ])
    metric_names = [
        ("TP (benar mendeteksi positif)", m["TP"]),
        ("FN (positif terlewat)", m["FN"]),
        ("FP (negatif keliru positif)", m["FP"]),
        ("TN (benar mendeteksi negatif)", m["TN"]),
        ("Recall", m["Recall"]), ("Accuracy", m["Accuracy"]),
        ("Precision", m["Precision"]), ("F1-score", m["F1-score"]),
        ("ROC-AUC", m["ROC-AUC"]),
        ("Accuracy Baseline Mayoritas", m["Accuracy Baseline Mayoritas"]),
        ("N Training", m["N Training"]), ("N Testing", m["N Testing"]),
    ]
    metrics_df = pd.DataFrame(metric_names, columns=["Indikator", "Nilai"])
    info_df = pd.DataFrame(list(cfg.items()), columns=["Parameter", "Nilai"])
    guide_df = pd.DataFrame([
        {"Nama Kolom": col,
         "Peran": "Target (Y)" if col == cfg["Target Y"] else "Identitas" if col in ID_COLUMNS else "Prediktor (X)" if col in result["X_train"].columns else "Tidak dipakai",
         "Deskripsi": GUIDE.get(col, "Variabel dari dataset unggahan pengguna")}
        for col in result["df"].columns
    ])
    cm_df = pd.DataFrame(result["cm"],
                         index=[f"Aktual: {neg}", f"Aktual: {pos}"],
                         columns=[f"Prediksi: {neg}", f"Prediksi: {pos}"])
    if profile is None:
        profile_df = pd.DataFrame({"Informasi": ["Status", "Keterangan"], "Nilai": [
            "Belum ada simulasi", "Jalankan prediksi profil pada Tab 4 untuk mengisi sheet ini."]})
    else:
        profile_df = pd.DataFrame(
            [("Input - " + str(k), v) for k, v in profile["input"].items()] + [
                ("Probabilitas kelas positif", profile["probability"]),
                ("Kelas positif", pos), ("Kelas prediksi", profile["prediction"]),
                ("Threshold", profile["threshold"]),
            ], columns=["Informasi", "Nilai"])
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_frames = [
            ("Pertanyaan & Jawaban", questions_df),
            ("Pengaturan Model", info_df),
            ("Data Lengkap", result["full"]),
            ("Data Training", result["training"]),
            ("Data Testing", result["testing"]),
            ("Prediksi Testing", result["prediction"]),
            ("Metrik Evaluasi", metrics_df),
            ("Simulasi Profil", profile_df),
            ("Kamus Variabel", guide_df),
        ]
        for sheet_name, frame in sheet_frames:
            # Perlakukan teks awalan "=" sebagai string, bukan rumus Excel.
            safe = frame.copy().astype(object)
            for col in safe.select_dtypes(include=["object"]).columns:
                safe[col] = safe[col].map(lambda v: "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
            safe.to_excel(writer, sheet_name=sheet_name, index=False)
        cm_df.to_excel(writer, sheet_name="Confusion Matrix", index=True)
        workbook = writer.book
        # Sheet pembuka: singkat, langsung menampilkan empat jawaban inti.
        cover = workbook.create_sheet("Ringkasan", 0)
        cover.merge_cells("A1:E2")
        cover["A1"] = "STATCAL ONLINE | REGRESI LOGISTIK"
        cover["A3"] = "Empat pertanyaan praktikum — DATA TESTING"
        cover["A5"], cover["B5"] = "Q1 • TP (benar terdeteksi)", m["TP"]
        cover["A6"], cover["B6"] = "Q2 • FN (risiko terlewat)", m["FN"]
        cover["A7"], cover["B7"] = "Q3 • Recall", m["Recall"]
        cover["A8"], cover["B8"] = "Q4 • Probabilitas profil", profile["probability"] if profile else "Belum dijalankan"
        cover["A10"], cover["B10"] = "Prediksi profil", profile["prediction"] if profile else "Belum tersedia"
        cover["A12"] = "Model: Logistic Regression | Dataset praktik simulasi / unggahan pengguna"
        cover["A13"] = f"Kelas positif (1): {pos}; kelas negatif (0): {neg}; threshold: {cfg['Threshold']:.2f}"
        cover["A14"] = f"Training: {m['N Training']} observasi | Testing: {m['N Testing']} observasi"
        cover["A16"] = "Semua TP, FN, recall, dan metrik evaluasi dihitung pada data TESTING."
        cover["A17"] = "Data Training adalah observasi yang digunakan melatih model; Data Testing adalah data evaluasi."
        cover["A18"] = "Probabilitas profil merupakan keluaran model, bukan kepastian atau hubungan sebab-akibat."
        cover["B7"].number_format = "0.0%"
        if profile:
            cover["B8"].number_format = "0.0%"
        cover["A1"].font = Font(size=17, bold=True, color="FFFFFF")
        cover["A1"].fill = PatternFill("solid", fgColor="11315C")
        cover.row_dimensions[1].height = 30
        cover.row_dimensions[2].height = 22
        cover.column_dimensions["A"].width = 39
        cover.column_dimensions["B"].width = 22
        cover.column_dimensions["C"].width = 12
        cover.column_dimensions["D"].width = 12
        cover.column_dimensions["E"].width = 12
        for r in [5, 6, 7, 8]:
            cover[f"A{r}"].font = Font(bold=True, color="11315C")
            cover[f"B{r}"].font = Font(bold=True, color="075A83", size=13)
            cover[f"A{r}"].fill = PatternFill("solid", fgColor="EAF4FC")
        for r in [16, 17, 18]:
            cover[f"A{r}"].font = Font(color="596C7D", italic=True, size=10)
        cover.sheet_view.showGridLines = False
        cover.sheet_properties.pageSetUpPr.fitToPage = True
        cover.page_setup.fitToWidth = 1
        cover.print_options.horizontalCentered = True
        for ws in workbook.worksheets:
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.orientation = "landscape" if ws.max_column >= 5 else "portrait"
            ws.sheet_view.zoomScale = 90
            if ws.title == "Ringkasan":
                ws.page_setup.orientation = "portrait"
                continue
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.print_title_rows = "1:1"
            ws.row_dimensions[1].height = 29
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill("solid", fgColor="143F65")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                col_letter = get_column_letter(col_idx)
                header = str(ws.cell(1, col_idx).value or "")
                sample_len = max((len(str(c.value)) for c in list(col_cells)[1:40] if c.value is not None), default=8)
                width = min(52, max(15, min(sample_len + 3, 42), len(header)+2))
                ws.column_dimensions[col_letter].width = width
            if ws.title in ("Pertanyaan & Jawaban", "Kamus Variabel"):
                ws.column_dimensions["B"].width = 50 if ws.title == "Pertanyaan & Jawaban" else 19
                ws.column_dimensions["E"].width = 62 if ws.title == "Pertanyaan & Jawaban" else ws.column_dimensions["E"].width
                if ws.title == "Kamus Variabel":
                    ws.column_dimensions["C"].width = 60
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                    ws.row_dimensions[row[0].row].height = 46 if ws.title == "Pertanyaan & Jawaban" else 32
            if ws.title in ("Data Testing", "Prediksi Testing"):
                for col_i in range(1, ws.max_column + 1):
                    if ws.cell(1, col_i).value in ("Probabilitas_Kelas_Positif", "Probabilitas_Kelas_Negatif"):
                        for cell in ws.iter_cols(min_col=col_i, max_col=col_i, min_row=2):
                            for c in cell: c.number_format = "0.00%"
            if ws.title == "Metrik Evaluasi":
                for row in range(2, ws.max_row + 1):
                    if ws.cell(row, 1).value in ("Recall", "Accuracy", "Precision", "F1-score", "ROC-AUC", "Accuracy Baseline Mayoritas"):
                        ws.cell(row, 2).number_format = "0.00%"
            if ws.title == "Pertanyaan & Jawaban":
                ws["D4"].number_format = "0.0%" # baris Q3
                if profile:
                    ws["D5"].number_format = "0.0%" # baris Q4
            if ws.title == "Simulasi Profil":
                for r in range(2, ws.max_row + 1):
                    if ws.cell(r, 1).value == "Probabilitas kelas positif":
                        ws.cell(r, 2).number_format = "0.0%"
    return output.getvalue()
